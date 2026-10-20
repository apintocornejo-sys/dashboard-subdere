"""
Asigna eventos según la columna "Fuente Financiamiento" (columna M) de la planilla
"Cartera PMU Catástrofe Región de Coquimbo 2026":

    Fuente Financiamiento                          ->  Evento del dashboard
    Transferencia Gobierno Regional 2026           ->  Transferencia Gobierno Regional 2026
    Ficha Simplificada SUBDERE Catastrofe 2026     ->  Transferencia SUBDERE Catastrofe 2026

- El evento antiguo "Ficha Simplificada SUBDERE Catastrofe 2026" se RENOMBRA a "Transferencia SUBDERE Catastrofe 2026"
  (orden del 20-oct-2026): todos los proyectos que lo tenían pasan al nombre nuevo, aunque no estén en la planilla.

- Cruza por "Id. Proyecto". Los proyectos de la planilla que no están en el dashboard se listan y no se tocan.
- Si un proyecto ya tenía OTRO evento, se REEMPLAZA (se pidió "incorporar"); el resumen lo lista antes de confirmar.
- Crea el evento nuevo en el desplegable de docs/index.html (hay que publicar con git después).
- Antes de escribir guarda respaldos (Datos adicionales y dashboard) en data_raw/. Se puede repetir sin problema.

Uso (desde la carpeta del proyecto):
    python3 scripts/asignar_evento_cartera_catastrofe.py --archivo RUTA/Cartera_PMU_Catastrofe_Region_de_Coquimbo_2026.xlsx --simular
    python3 scripts/asignar_evento_cartera_catastrofe.py --archivo RUTA/Cartera_PMU_...xlsx
Después: git add -A && git commit -m "Evento Transferencia SUBDERE Catastrofe 2026" && git push
"""
import re
import sys
import json
import shutil
import argparse
import collections
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import agregar_evento_dashboard as agr
import asignar_evento_listado as lst

BASE_DIR = lst.BASE_DIR
INDEX_HTML = BASE_DIR / "docs" / "index.html"
FUENTE_A_EVENTO = {
    "Transferencia Gobierno Regional 2026": "Transferencia Gobierno Regional 2026",
    "Ficha Simplificada SUBDERE Catastrofe 2026": "Transferencia SUBDERE Catastrofe 2026",
}
EVENTO_NUEVO = "Transferencia SUBDERE Catastrofe 2026"
EVENTO_ANTIGUO = "Ficha Simplificada SUBDERE Catastrofe 2026"      # se renombra a EVENTO_NUEVO
_FUENTES_NORM = {lst.normalizar(k): v for k, v in FUENTE_A_EVENTO.items()}


def id_limpio(i) -> str:
    s = re.sub(r"\s+", "", str(i or "")).upper()
    return re.sub(r"\.0$", "", s)


def leer_cartera(ruta) -> dict:
    """{id_proyecto: fuente_de_financiamiento} desde la columna 'Fuente Financiamiento' (columna M) de la planilla."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True, read_only=True)
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        for k, fila in enumerate(filas):
            nombres = [lst.normalizar(c) for c in fila]
            if "id proyecto" in nombres and "fuente financiamiento" in nombres:
                ci, cf = nombres.index("id proyecto"), nombres.index("fuente financiamiento")
                out = {}
                for r in filas[k + 1:]:
                    if len(r) > max(ci, cf) and r[ci] and str(r[cf] or "").strip():
                        out[id_limpio(r[ci])] = str(r[cf]).strip()
                return out
    raise ValueError("No encontré las columnas 'Id. Proyecto' y 'Fuente Financiamiento' en la planilla.")


def planificar(cartera: dict, ids_dashboard: set, datos_extra: dict) -> dict:
    """Calcula (sin escribir) qué cambia. Pura: se prueba sin Drive."""
    ids_norm = {id_limpio(i): i for i in ids_dashboard}      # id de la planilla -> id real del dashboard
    nuevos, reemplazos, ya, sin_regla, fuera = {}, {}, 0, collections.Counter(), []
    for idp, fuente in cartera.items():
        evento = _FUENTES_NORM.get(lst.normalizar(fuente))
        if not evento:
            sin_regla[fuente] += 1
            continue
        real = ids_norm.get(idp)
        if not real:
            fuera.append(idp)
            continue
        actual = ((datos_extra.get(real) or {}).get("evento") or "").strip()
        if actual == evento:
            ya += 1
        elif actual:
            reemplazos[real] = (evento, actual)
        else:
            nuevos[real] = evento
    # Renombre del evento antiguo: todo proyecto que lo tenga (esté o no en la planilla) pasa al nombre nuevo.
    asignados = set(nuevos) | set(reemplazos)
    renombrar = {i: EVENTO_NUEVO for i, d in datos_extra.items()
                 if (d or {}).get("evento", "").strip() == EVENTO_ANTIGUO and i not in asignados}
    return {"nuevos": nuevos, "reemplazos": reemplazos, "ya": ya, "sin_regla": sin_regla, "fuera": fuera, "renombrar": renombrar}


def aplicar(datos_extra: dict, asignar: dict) -> dict:
    out = json.loads(json.dumps(datos_extra))
    for idp, evento in asignar.items():
        e = out.get(idp) or {}
        e["evento"] = evento
        out[idp] = e
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--archivo", required=True, help="Excel de la cartera (Cartera_PMU_Catastrofe_...xlsx)")
    ap.add_argument("--simular", action="store_true")
    a = ap.parse_args(argv)
    if not INDEX_HTML.exists():
        sys.exit(f"No encuentro {INDEX_HTML}. Ejecuta esto desde la carpeta del proyecto.")
    cartera = leer_cartera(a.archivo)
    servicio = lst.obtener_servicio_drive()
    proyectos = lst.cargar_proyectos(servicio)
    datos_extra = lst.leer_json_drive(servicio, lst.DATOS_EXTRA_FILE_ID)
    html = INDEX_HTML.read_text(encoding="utf-8")
    existentes = agr.leer_eventos(agr.RE_BLOQUE.search(html).group(2))
    plan = planificar(cartera, {p["id_proyecto"] for p in proyectos}, datos_extra)
    asignar = {**plan["nuevos"], **{i: v[0] for i, v in plan["reemplazos"].items()}, **plan["renombrar"]}
    crear_evento = EVENTO_NUEVO.lower() not in (e.lower() for e in existentes) and any(e == EVENTO_NUEVO for e in asignar.values())
    quitar_antiguo = EVENTO_ANTIGUO in existentes
    por_evento = collections.Counter(asignar.values())
    nombres = {p["id_proyecto"]: p.get("nombre_proyecto", "") for p in proyectos}

    print(f"Planilla: {len(cartera)} proyectos con fuente de financiamiento · dashboard: {len(proyectos)} proyectos\n")
    print(f"Ya tenían su evento: {plan['ya']}   ·   se asignan (sin evento): {len(plan['nuevos'])}   ·   se REEMPLAZA otro evento: {len(plan['reemplazos'])}")
    for ev, n in por_evento.items():
        print(f"   {n:>4}  -> {ev}")
    print(f"Renombrados desde '{EVENTO_ANTIGUO}': {len(plan['renombrar'])} proyectos (no estaban en la planilla)")
    if crear_evento:
        print(f"\nSe CREARÁ el evento '{EVENTO_NUEVO}' en el dashboard.")
    if quitar_antiguo:
        print(f"Se ELIMINARÁ del desplegable el evento antiguo '{EVENTO_ANTIGUO}'.")
    if plan["reemplazos"]:
        print("\nSe REEMPLAZA el evento de:")
        for idp, (nuevo, antes) in sorted(plan["reemplazos"].items()):
            print(f"   {idp} | {str(nombres.get(idp))[:50]} | antes: {antes} -> {nuevo}")
    if plan["fuera"]:
        print(f"\nEn la planilla pero NO están en el dashboard ({len(plan['fuera'])}; no se tocan): " + ", ".join(plan["fuera"][:30]))
    for f, n in plan["sin_regla"].items():
        print(f"\nFuente sin regla (no se toca): '{f}' · {n} proyectos")
    if a.simular:
        print("\n(Simulación: no se escribió nada.)")
        return
    if not asignar and not crear_evento and not quitar_antiguo:
        print("\nNada que hacer: todo está al día.")
        return
    if input(f"\n¿Aplicar? Se asigna evento a {len(asignar)} proyectos en Drive" + (" y se crea el evento en docs/index.html" if crear_evento else "") + " [s/N] ").strip().lower() not in ("s", "si", "sí"):
        print("No se aplicó nada.")
        return
    lst.RESPALDOS_DIR.mkdir(exist_ok=True)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    resp = lst.RESPALDOS_DIR / f"datos_extra_respaldo_{marca}.json"
    resp.write_text(json.dumps(datos_extra, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Respaldo de Datos adicionales: {resp.name}")
    if crear_evento or quitar_antiguo:     # primero el dashboard: si falla la verificación, no se escribe nada en Drive
        import asignar_evento_planilla as pl
        html_nuevo = html
        if quitar_antiguo:      # renombra en la misma posición (o solo quita el antiguo si el nuevo ya existe)
            html_nuevo, msg = pl.reemplazar_evento_html(html_nuevo, EVENTO_ANTIGUO, EVENTO_NUEVO)
            print(msg)
        else:
            html_nuevo, msg = agr.agregar(html_nuevo, EVENTO_NUEVO)
            print(msg)
        shutil.copy2(INDEX_HTML, lst.RESPALDOS_DIR / f"index.html.bak_evento_{marca}")
        INDEX_HTML.write_text(html_nuevo, encoding="utf-8")
    lst.escribir_json_drive(servicio, lst.DATOS_EXTRA_FILE_ID, aplicar(datos_extra, asignar))
    print(f"Listo: {len(asignar)} proyectos con evento asignado.")
    if crear_evento or quitar_antiguo:
        print('Ahora publica el evento nuevo: git add -A && git commit -m "Evento Transferencia SUBDERE Catastrofe 2026" && git push')


if __name__ == "__main__":
    main()
